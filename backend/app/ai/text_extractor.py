"""Extracts plain text from uploaded files so the AI modules can analyze them (in memory only)."""
from __future__ import annotations

import io
from pathlib import Path

MAX_CHARS = 200_000  # enough for analysis, avoids huge memory use


def extract_text(file_bytes: bytes, filename: str) -> str:
    extension = Path(filename).suffix.lower()
    try:
        if extension in {".txt", ".md", ".csv"}:
            text = _decode_plain_text(file_bytes)
        elif extension == ".pdf":
            text = _extract_pdf(file_bytes)
        elif extension == ".docx":
            text = _extract_docx(file_bytes)
        else:
            text = ""
    except Exception:  # corrupted or unsupported content: analyze nothing rather than crash
        text = ""
    return text[:MAX_CHARS]


def _decode_plain_text(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


def _extract_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(file_bytes: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(file_bytes))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
